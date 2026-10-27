"""Import cached Excel values into a private seed; never re-evaluate external formulas."""
import argparse
import datetime
import json
import os
from pathlib import Path

import openpyxl

parser = argparse.ArgumentParser(description='엑셀을 새 설치의 초기 데이터로 변환합니다. 기존 DB는 변경하지 않습니다.')
parser.add_argument('excel', type=Path)
parser.add_argument('output', type=Path)
args = parser.parse_args()
if args.output.exists():
    parser.error('출력 파일이 이미 있습니다. 기존 데이터를 보호하기 위해 다른 경로를 지정하세요.')
workbook = openpyxl.load_workbook(args.excel, data_only=True)
sheet = workbook['통합_자산현황_최종']
holdings = []
for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True)):
    owner, account, ticker, name, quantity, price, _, value, rate, *_ = row
    if not owner:
        continue
    if not all(isinstance(v, (int, float)) for v in (quantity, price, value, rate)):
        raise ValueError('자산현황 계산 결과가 없습니다. Google Sheets에서 다시 Excel로 내려받아 주세요.')
    currency = 'KRW' if ticker == 'CASH' or ticker.startswith('KRX:') else 'USD'
    holdings.append(dict(id=f'h{index+1}', owner=owner, account=account, ticker=ticker.upper(), name=name,
                         quantity=quantity, price=price, currency=currency,
                         costPrice=price / (1 + rate) if rate and rate != -1 else price))
transactions = []
for index, row in enumerate(workbook['통합_거래내역'].iter_rows(min_row=2, values_only=True)):
    if not row[0]:
        continue
    kind, owner, account, ticker, name, quantity, krw, usd, amount, *_ = row
    transactions.append(dict(id=f'import-{index+1}', type=kind, owner=owner, account=account,
                             ticker=ticker.upper(), name=name, quantity=quantity, price=usd or krw or 0,
                             currency='USD' if usd else 'KRW', amount=amount, date=None, imported=True))
history = [dict(date=r[0].date().isoformat(), value=r[1])
           for r in workbook['자산추이'].iter_rows(min_row=2, values_only=True)
           if isinstance(r[0], datetime.datetime) and isinstance(r[1], (int, float))]
owners = list(dict.fromkeys(h['owner'] for h in holdings))
fx = sheet['M1'].value
if len(owners) != 2 or not isinstance(fx, (int, float)) or fx <= 0:
    raise ValueError('두 소유자 및 원/달러 환율을 확인해 주세요.')
seed = dict(version=0, fx=fx, holdings=holdings, transactions=transactions, history=history,
            quoteUpdatedAt=None, source=args.excel.name, owners=owners)
args.output.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
fd = os.open(args.output, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
with os.fdopen(fd, 'w') as stream:
    json.dump(seed, stream, ensure_ascii=False, indent=2)
print(f'가져오기 완료: 자산 {len(holdings)}개 / 거래 {len(transactions)}건 / 추이 {len(history)}개')
print('기존 거래를 재계산하지 않고 자산현황을 초기 잔고로 사용합니다. 기존 DB가 있으면 자동으로 덮어쓰지 않습니다.')
